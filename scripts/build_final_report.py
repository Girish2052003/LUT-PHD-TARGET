#!/usr/bin/env python3
"""Build the editable report from its reviewed Markdown source (no model calls).

Requires python-docx. PDF conversion and page-image review are separate steps;
this builder does not claim that successful generation proves visual quality.
"""
from pathlib import Path
import argparse
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TOKEN = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def inline(paragraph, text, size=None):
    for token in TOKEN.split(text):
        if not token:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
        if match:
            link = OxmlElement("w:hyperlink")
            rid = paragraph.part.relate_to(
                match[2], "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True)
            link.set(qn("r:id"), rid)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), "254D72")
            props.append(color)
            if size:
                s = OxmlElement("w:sz"); s.set(qn("w:val"), str(int(size * 2))); props.append(s)
            run.append(props)
            content = OxmlElement("w:t"); content.text = match[1]; run.append(content)
            link.append(run); paragraph._p.append(link)
            continue
        bold = token.startswith("**")
        code = token.startswith("`")
        italic = token.startswith("*") and not bold
        value = token[2:-2] if bold else token[1:-1] if code or italic else token
        run = paragraph.add_run(value)
        run.bold = bold; run.italic = italic
        if code:
            run.font.name = "DejaVu Sans Mono"; run.font.size = Pt(9)
        elif size:
            run.font.size = Pt(size)


def field(paragraph, name):
    node = OxmlElement("w:fldSimple"); node.set(qn("w:instr"), name)
    paragraph._p.append(node)


def build(source, output):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
    sec.header_distance = Inches(.27); sec.footer_distance = Inches(.28)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Serif"; normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Liberation Sans"; style.font.size = Pt(size)
        style.font.bold = True; style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        # Remove any decorative rules inherited from the runtime's base template.
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    header = sec.header.paragraphs[0]
    header.add_run("SPDZCoder research assignment • Girish Nallan Chakravathy").font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Draft for author review • 22 October 2026    |    ").font.size = Pt(8)
    field(footer, "PAGE"); footer.add_run(" / "); field(footer, "NUMPAGES")
    doc.core_properties.title = "SPDZCoder Reproduction and Privacy Modification"
    doc.core_properties.author = "Girish Nallan Chakravathy"
    doc.core_properties.subject = "Reproduction, scoped privacy modification and measured comparison"
    lines = source.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1; continue
        if line == "<!-- PAGEBREAK -->":
            doc.add_page_break(); i += 1; continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], "Title"); i += 1; continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1"); i += 1; continue
        if line.startswith("```"):
            i += 1; code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i]); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(7)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_together = True
            run = p.add_run("\n".join(code)); run.font.name = "DejaVu Sans Mono"; run.font.size = Pt(8.5)
            i += 1; continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [s.strip() for s in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.autofit = False
            widths = {3: [4.1, 1.45, 1.45], 4: [4.15, .95, .95, .95], 5: [2.65, .9, 1.05, 1.3, 1.1]}[len(rows[0])]
            for column, width in zip(table.columns, widths): column.width = Inches(width)
            props = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                e = OxmlElement("w:" + edge); e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "C3C8CC"); borders.append(e)
            props.append(borders)
            for r, values in enumerate(rows):
                row = table.add_row()
                trprops = row._tr.get_or_add_trPr()
                no_split = OxmlElement("w:cantSplit"); trprops.append(no_split)
                if r == 0:
                    repeat = OxmlElement("w:tblHeader"); trprops.append(repeat)
                for c, value in enumerate(values):
                    cell = row.cells[c]; cell.width = Inches(widths[c])
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1
                    inline(p, value, 10)
                    if c: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    if r == 0:
                        for run in p.runs: run.bold = True
                        shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "EFF1F3")
                        cell._tc.get_or_add_tcPr().append(shading)
            continue
        # Blank lines delimit paragraphs. Explicit Markdown hard breaks stay visible.
        parts = []
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "```", "<!--")):
            parts.append(lines[i]); i += 1
        p = doc.add_paragraph()
        text = "\n".join(x.rstrip() for x in parts) if any(x.endswith("  ") for x in parts) else " ".join(parts)
        ref = re.match(r"\[\d+\] ", text)
        inline(p, text, 9 if ref else None)
        if ref:
            p.paragraph_format.space_after = Pt(3)
        if text.startswith("*Table "):
            p.paragraph_format.space_before = Pt(3)
            for run in p.runs: run.font.size = Pt(9)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Created {output.relative_to(ROOT)} from {source.relative_to(ROOT)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "docs/final_report.md")
    parser.add_argument("--output", type=Path, default=ROOT / "docs/submission/SPDZCoder_Report_Girish.docx")
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
